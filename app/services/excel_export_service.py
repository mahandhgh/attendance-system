from pathlib import Path
from datetime import timedelta
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter


class ExcelExportService:

    @staticmethod
    def timedelta_to_excel_time(value: timedelta) -> float:
        if value is None:
            return 0.0

        return value.total_seconds() / 86400

    @staticmethod
    def decimal_to_float(value):
        if value is None:
            return 0.0

        if isinstance(value, Decimal):
            return float(value)

        return float(value)

    def export(self, final_reports, output_path):
        output_path = Path(output_path)

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Monthly Reports"

        headers = [
            "Employee ID",
            "Jalali Year",
            "Jalali Month",

            "Required Time",
            "Worked Time",
            "Overtime",
            "Shortage",
            "Delay",

            "Leave Time",
            "Allowed Leave",
            "Excess Leave",

            "Work Days",
            "Absent Days",
            "Incomplete Days",
            "Multiple Punch Days",
            "Holiday Days",
            "Leave Days",

            "Hourly Rate",
            "Overtime Pay",
            "Shortage Deduction",
            "Excess Leave Deduction",
            "Total Deduction",
            "Net Adjustment",
        ]

        worksheet.append(headers)

        for result in final_reports:
            report = result.attendance_report
            salary = result.salary_calculation

            worksheet.append([
                report.employee_id,
                result.year,
                result.month,

                self.timedelta_to_excel_time(
                    report.total_required_time
                ),

                self.timedelta_to_excel_time(
                    report.total_worked_time
                ),

                self.timedelta_to_excel_time(
                    report.total_overtime
                ),

                self.timedelta_to_excel_time(
                    report.total_shortage
                ),

                self.timedelta_to_excel_time(
                    report.total_delay
                ),

                self.timedelta_to_excel_time(
                    report.total_leave_time
                ),

                self.timedelta_to_excel_time(
                    report.total_allowed_leave_time
                ),

                self.timedelta_to_excel_time(
                    report.total_excess_leave_time
                ),

                report.work_days,
                report.absent_days,
                report.incomplete_days,
                report.multiple_punch_days,
                report.holiday_days,
                report.leave_days,

                self.decimal_to_float(
                    salary.hourly_rate
                ),

                self.decimal_to_float(
                    salary.overtime_pay
                ),

                self.decimal_to_float(
                    salary.shortage_deduction
                ),

                self.decimal_to_float(
                    salary.excess_leave_deduction
                ),

                self.decimal_to_float(
                    salary.total_deduction
                ),

                self.decimal_to_float(
                    salary.net_adjustment
                ),
            ])

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        worksheet.sheet_view.showGridLines = False

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="0F766E",
        )

        header_font = Font(
            bold=True,
            color="FFFFFF",
        )

        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        thin_side = Side(
            style="thin",
            color="D1D5DB",
        )

        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for cell in worksheet[1]:

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border

        worksheet.row_dimensions[1].height = 32

        time_columns = {
            "D",  # Required
            "E",  # Worked
            "F",  # Overtime
            "G",  # Shortage
            "H",  # Delay
            "I",  # Leave
            "J",  # Allowed Leave
            "K",  # Excess Leave
        }

        money_columns = {
            "R",  # Hourly Rate
            "S",  # Overtime Pay
            "T",  # Shortage Deduction
            "U",  # Excess Leave Deduction
            "V",  # Total Deduction
            "W",  # Net Adjustment
        }

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
        ):

            for cell in row:

                cell.border = border

                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="center",
                )

            for column in time_columns:
                worksheet[f"{column}{row[0].row}"].number_format = (
                    "[h]:mm:ss"
                )

            for column in money_columns:
                worksheet[f"{column}{row[0].row}"].number_format = (
                    "#,##0.00"
                )

        if worksheet.max_row >= 2:
            table_ref = (
                f"A1:W{worksheet.max_row}"
            )

            table = Table(
                displayName="MonthlyReportsTable",
                ref=table_ref,
            )

            table_style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = table_style

            worksheet.add_table(table)

        worksheet.conditional_formatting.add(
            f"W2:W{worksheet.max_row}",
            CellIsRule(
                operator="lessThan",
                formula=["0"],
                fill=PatternFill(
                    fill_type="solid",
                    fgColor="FEE2E2",
                ),
                font=Font(
                    color="B91C1C",
                    bold=True,
                ),
            ),
        )

        worksheet.conditional_formatting.add(
            f"W2:W{worksheet.max_row}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill(
                    fill_type="solid",
                    fgColor="DCFCE7",
                ),
                font=Font(
                    color="15803D",
                    bold=True,
                ),
            ),
        )

        column_widths = {
            "A": 18,
            "B": 12,
            "C": 14,

            "D": 17,
            "E": 17,
            "F": 17,
            "G": 17,
            "H": 17,

            "I": 17,
            "J": 17,
            "K": 17,

            "L": 12,
            "M": 12,
            "N": 16,
            "O": 20,
            "P": 14,
            "Q": 12,

            "R": 17,
            "S": 17,
            "T": 21,
            "U": 24,
            "V": 18,
            "W": 18,
        }

        for column, width in column_widths.items():

            worksheet.column_dimensions[
                column
            ].width = width

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.row_dimensions[
                row_number
            ].height = 22

        worksheet.page_setup.orientation = (
            "landscape"
        )

        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

        worksheet.sheet_properties.pageSetUpPr.fitToPage = True

        worksheet.print_title_rows = "1:1"

        workbook.save(output_path)